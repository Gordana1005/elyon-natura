# -*- coding: utf-8 -*-
"""Stock v2 — the owner's review workbook (Macedonian): exports/stock/Magacin-pregled-<date>.xlsx.

Reads the outputs of docs/stock/build_sigma_stock.py (exports/stock/*.json) and lays them out for the owner:
  Резиме · Почетна состојба · Набавни цени · Производи→Артикли · collabBox шифри без артикл · Web производи ·
  Пакети · Сигма документи од 22.09
No customer data: articles, products, Sigma documents, company clients and Sigma user logins only. The workbook
holds purchase costs (owners only) — it stays in the gitignored exports/ folder and is never committed.

  python docs/stock/mapping_review_xlsx.py [--out DIR] [--date 2026-10-01]
"""
import argparse, collections, datetime, json, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))
ap = argparse.ArgumentParser()
ap.add_argument("--out", default=os.path.join(ROOT, "exports", "stock"))
ap.add_argument("--date", default=datetime.date.today().isoformat())
ARGS = ap.parse_args()
OUT = os.path.abspath(ARGS.out)


def load(name):
    with open(os.path.join(OUT, name), encoding="utf-8") as f:
        return json.load(f)


SUMMARY = load("summary.json")
ARTICLES = {a["code"]: a for a in load("articles.json")}
LOCAL = {a["code"]: a for a in load("local-articles.json")}
COSTS = {c["code"]: c for c in load("costs.json")}
COST_REVIEW = {c["code"]: c for c in load("review-costs.json")}
OPENINGS = load("openings.json")
OPEN_REVIEW = load("review-openings.json")
RECIPES = load("review-recipes.json")
CB = load("review-collabbox-codes.json")
WEB = load("review-web.json")
KITS = load("review-kits.json")
DOCS = load("review-sigma-docs.json")
EXEMPT = load("exempt.json")
MKD_PER_EUR = 61.5

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLE = Font(name="Arial", size=14, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="2F5D50")
TOTAL_FILL = PatternFill("solid", fgColor="E8F0EC")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
THIN = Side(style="thin", color="C9D3CF")
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
UNITS = "#,##0"
UNITS3 = "#,##0.###"
MONEY = "#,##0.00"
CONF_MK = {"high": "висока", "medium": "средна", "low": "ниска", None: "—"}
BASIS_MK = {
    "04_calcbuy_qtyweighted": "04 CalcBuyPrice, просек по количина",
    "04_calcbuy_single": "04 CalcBuyPrice",
    "04_last_buyprice_2026": "04 набавна на продажбите 2026 (медијана)",
    "other_object": "друг објект на Ф00001",
    "none": "нема цена",
}
FLAG_MK = {"buckets_differ": "различни цени по корпа", "cost_zero": "залиха со цена 0", "no_stock_04": "нема залиха во 04",
           "other_object": "цена од друг објект"}


def article_name(code):
    return (ARTICLES.get(code) or LOCAL.get(code) or {}).get("name", "")


def recipe_text(lines):
    if not lines:
        return ""
    role = {"gift": " (подарок)", "component": "", "main": ""}
    return " + ".join(f"{l['qty']:g} × {l['code']} {article_name(l['code'])}{role.get(l.get('role'), '')}" for l in lines)


def sheet(wb, title, headers, rows, widths, formats=None, freeze_cols=1, wrap_cols=(), total_row=None, note=None):
    """A table sheet: optional note on row 1, header, optional totals row, data; Arial, frozen header, widths."""
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(name="Arial", size=10, italic=True)
        ws.row_dimensions[1].height = 30
        ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(headers), 10))
        r0 = 2
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r0].height = 32
    start = r0 + 1
    if total_row is not None:
        for c, v in enumerate(total_row, 1):
            cell = ws.cell(row=start, column=c, value=v)
            cell.font, cell.fill = BOLD, TOTAL_FILL
            if formats and formats.get(c):
                cell.number_format = formats[c]
        start += 1
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=start + i, column=c, value=v)
            cell.font = FONT
            cell.alignment = WRAP if c in wrap_cols else TOP
            if formats and formats.get(c):
                cell.number_format = formats[c]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1 + (1 if total_row is not None else 0), column=freeze_cols + 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{max(start + len(rows) - 1, r0)}"
    return ws, start


wb = Workbook()
wb.remove(wb.active)

# ── Резиме ──────────────────────────────────────────────────────────────────────────────────────────────────
ws = wb.create_sheet("Резиме")
S = SUMMARY
v = OPENINGS["variants"]
lines = [
    ("Магацин — преглед за сопственикот", None, TITLE),
    (f"Направено {ARGS.date} од Сигма извозот ({S['sigma_export'][:16].replace('T', ' ')}) и CRM (само читање). "
     "Ништо не е запишано во CRM — ова е предлог за одобрување.", None, FONT),
    ("", None, FONT),
    ("Почетна состојба на 22.09 (готови производи, парчиња)", None, BOLD),
    ("Сигма 04 — утро 22.09 (документи до 21.09)", v["04_morning_2209"]["totals"]["units_fg_incl_negatives"], FONT),
    ("Сигма 04 — крај 22.09 (вклучен е ТН1 04-00045 04→08)", v["04_end_2209"]["totals"]["units_fg_incl_negatives"], FONT),
    ("Сигма 04 + 08 — утро 22.09", v["04_plus_08_morning_2209"]["totals"]["units_fg_incl_negatives"], FONT),
    ("Сигма 08 — крај 22.09", v["08_end_2209"]["totals"]["units_fg_incl_negatives"], FONT),
    ("Вредност 04 утро по Сигма набавна (ден)", v["04_morning_2209"]["totals"]["value_mkd"], FONT),
    ("Пописната листа од 22.09 се внесува во листот „Почетна состојба“, жолтата колона.", None, FONT),
    ("", None, FONT),
    ("Артикли (Сигма шифри)", None, BOLD),
    ("Артикли вкупно (АРТИКЛ, ТС, ЛОЈАЛИТИ)", S["articles"]["total"], FONT),
    ("Активни (продадени / испратени / на залиха)", S["articles"]["active"], FONT),
    ("Сетови / пакети како посебна шифра", S["articles"]["sets"], FONT),
    ("Нови локални артикли (ги нема во Сигма)", S["articles"]["local_proposed"], FONT),
    ("", None, FONT),
    ("Набавни цени (Сигма CalcBuyPrice, денари без ДДВ)", None, BOLD),
    ("Активни артикли со цена", f"{S['costs']['active_costed']} од {S['costs']['active_articles']}", FONT),
    ("Активни CRM производи со целосно пресметана набавна (по рецепт)", f"{S['recipes']['active_costed']} од {S['recipes']['active_products']}", FONT),
    ("", None, FONT),
    ("Производи → артикли (рецепти)", None, BOLD),
    ("Предлози вкупно", S["recipes"]["total"], FONT),
    ("— висока доверба (се одобруваат автоматски)", S["recipes"]["by_confidence"].get("high", 0), FONT),
    ("— средна (ги гледа сопственикот)", S["recipes"]["by_confidence"].get("medium", 0), FONT),
    ("— ниска (ги гледа сопственикот)", S["recipes"]["by_confidence"].get("low", 0), FONT),
    ("Активни производи без предлог", S["recipes"]["active_products"] - S["recipes"]["active_with_proposal"], FONT),
    ("", None, FONT),
    ("Сигма документи од 22.09 (04 / 08 / 11)", None, BOLD),
    ("Книжени документи", S["batch"]["docs"], FONT),
    ("— би влегле во Движења", S["batch"]["verdicts"].get("included", 0), FONT),
    ("— исклучени (MEX фактура, АД Астра, 04↔08 …)", S["batch"]["verdicts"].get("excluded", 0), FONT),
    ("Нацрти (најавено, не мрдаат залиха)", S["batch"]["drafts"], FONT),
    ("", None, FONT),
    ("Пратки од 22.09 (MEX)", None, BOLD),
]
for p in S.get("parcels_since_2209", []):
    lines.append((f"{p['account']}: пратки / со collabBox / web / CRM нарачка",
                  f"{p['parcels']:,} / {p['with_collabbox']:,} / {p['with_web']:,} / {p['with_order']:,}".replace(",", "."), FONT))
lines += [("", None, FONT), ("Како се чита", None, BOLD),
          ("Жолто = внесувате вие. Зелено = предлог со висока доверба. Црвено = проверете.", None, FONT),
          ("Цените и вредностите ги гледаат само сопствениците. Датотеката не се праќа надвор.", None, FONT)]
for i, (label, value, font) in enumerate(lines, 1):
    a = ws.cell(row=i, column=1, value=label)
    a.font = font
    if value is not None:
        b = ws.cell(row=i, column=2, value=value)
        b.font = FONT
        if isinstance(value, (int, float)):
            b.number_format = UNITS
        b.alignment = Alignment(horizontal="right")
ws.column_dimensions["A"].width = 78
ws.column_dimensions["B"].width = 26

# ── Почетна состојба ────────────────────────────────────────────────────────────────────────────────────────
KEYS = ["04_morning_2209", "04_end_2209", "04_plus_08_morning_2209", "08_end_2209", "04_now", "08_now"]
rows = []
codes = sorted(OPEN_REVIEW, key=lambda c: (-(OPEN_REVIEW[c].get("04_morning_2209") or 0), c))
for code in codes:
    o = OPEN_REVIEW[code]
    if not any(o.get(k) for k in KEYS[:4]):
        continue
    a = ARTICLES.get(code, {})
    cost = (COSTS.get(code) or {}).get("cost_mkd")
    m = o.get("04_morning_2209") or 0
    flags = []
    if any((o.get(k) or 0) < 0 for k in KEYS[:4]):
        flags.append("минус во Сигма")
    if a.get("unit") == "КОМ" and any(abs((o.get(k) or 0) - round(o.get(k) or 0)) > 1e-9 for k in KEYS[:4]):
        flags.append("дропка кај КОМ")
    if not m and (o.get("08_end_2209") or 0) > 0:
        flags.append("само во 08")
    rows.append([code, a.get("name"), a.get("unit"), a.get("sigma_class"), a.get("brand")] +
                [o.get(k) or 0 for k in KEYS] +
                [cost, (m * cost) if cost is not None else None, None, None, ", ".join(flags)])
headers = ["Шифра", "Артикл", "Ед.", "Класа", "Бренд", "04 утро 22.09", "04 крај 22.09", "04+08 утро 22.09", "08 крај 22.09",
           "Сигма 04 денес (30.09)", "Сигма 08 денес", "Набавна (ден)", "Вредност 04 утро (ден)",
           "ПОПИС 22.09 (внесете)", "Разлика попис − 04 утро", "Белешка"]
fmt = {6: UNITS3, 7: UNITS3, 8: UNITS3, 9: UNITS3, 10: UNITS3, 11: UNITS3, 12: MONEY, 13: UNITS, 14: UNITS3, 15: UNITS3}
first = 3 + 1  # note row 1, header row 2, totals row 3, data from 4
last = first + len(rows) - 1
total = ["", "ВКУПНО (сите артикли)", "", "", ""] + [f"=SUM({get_column_letter(c)}{first}:{get_column_letter(c)}{last})" for c in range(6, 12)] + \
        ["", f"=SUM(M{first}:M{last})", f"=SUM(N{first}:N{last})", f"=SUM(O{first}:O{last})", ""]
ws, start = sheet(wb, "Почетна состојба", headers, rows,
                  [9, 44, 6, 10, 20, 12, 12, 13, 12, 13, 12, 11, 15, 15, 14, 22], fmt, freeze_cols=2, wrap_cols=(2,),
                  total_row=total,
                  note="Состојба во Сигма по датум на документ (содржина како во извозот 30.09). Готовите производи во 04 утро на 22.09 = "
                       f"{v['04_morning_2209']['totals']['units_fg_incl_negatives']:,.0f}; крај на 22.09 = {v['04_end_2209']['totals']['units_fg_incl_negatives']:,.0f}; "
                       f"08 = {v['08_end_2209']['totals']['units_fg_incl_negatives']:,.0f} парчиња. Внесете ја пописната листа во жолтата колона.".replace(",", "."))
for i in range(len(rows)):
    r = start + i
    ws.cell(row=r, column=14).fill = INPUT_FILL
    ws.cell(row=r, column=15, value=f'=IF(N{r}="","",N{r}-F{r})').number_format = UNITS3
    ws.cell(row=r, column=15).font = FONT
    if rows[i][-1]:
        ws.cell(row=r, column=16).fill = WARN_FILL

# ── Набавни цени ────────────────────────────────────────────────────────────────────────────────────────────
old_by_code = collections.defaultdict(list)
for r in RECIPES:
    if r["lines"] and len(r["lines"]) == 1 and r["confidence"] in ("high", "medium") and r.get("cost_price_eur"):
        l = r["lines"][0]
        old_by_code[l["code"]].append((r["product_name"], float(r["cost_price_eur"]) * MKD_PER_EUR / (l["qty"] or 1)))
products_by_code = collections.defaultdict(list)
for r in RECIPES:
    for l in r["lines"] or []:
        products_by_code[l["code"]].append(r["product_name"])
rows = []
for code, c in sorted(COSTS.items(), key=lambda kv: (not ARTICLES.get(kv[0], {}).get("active"), kv[0])):
    a = ARTICLES.get(code, {})
    if not a.get("active") and code not in products_by_code:
        continue
    rv = COST_REVIEW.get(code, {})
    old = old_by_code.get(code)
    old_mkd = sum(x[1] for x in old) / len(old) if old else None
    buckets = "; ".join(f"{b['wyear']}: {b['qty']:,.0f} @ {b['price']:,.2f}" for b in rv.get("buckets_04", []))
    rows.append([code, a.get("name"), a.get("brand"), "да" if a.get("active") else "не", c["cost_mkd"], BASIS_MK.get(c["basis"], c["basis"]),
                 c.get("source_ref"), ", ".join(FLAG_MK.get(f, f) for f in c["flags"]), buckets.replace(",", "·"),
                 rv.get("sales_lines_2026"), rv.get("sales_price_2026_median"), old_mkd,
                 (c["cost_mkd"] / old_mkd) if (old_mkd and c["cost_mkd"]) else None,
                 "; ".join(sorted(set(products_by_code.get(code, []))))[:300]])
headers = ["Шифра", "Артикл", "Бренд", "Активен", "Набавна Сигма (ден)", "Основа", "Извор", "Ознаки", "Корпи во 04 (WYear: количина @ цена)",
           "Продажни ставки 2026", "Медијана набавна на продажбите 2026", "Стара CRM цена (ден = € × 61,5)", "Сигма / стара",
           "CRM производи"]
ws, start = sheet(wb, "Набавни цени", headers, rows, [9, 40, 18, 8, 13, 30, 30, 24, 40, 11, 14, 14, 10, 50],
                  {5: MONEY, 10: UNITS, 11: MONEY, 12: MONEY, 13: "0.00"}, freeze_cols=2, wrap_cols=(2, 9, 14),
                  note="Сигма CalcBuyPrice од Ф00001 (BioNatural по цена на производство, никогаш по цената на АД Астра), денари без ДДВ. "
                       "„Стара CRM цена“ = products.cost_price (€) × 61,5 за производите што се еден артикл. Се применува за целата историја.")
for i, row in enumerate(rows):
    if row[4] is None:
        ws.cell(row=start + i, column=5).fill = WARN_FILL
    if row[7]:
        ws.cell(row=start + i, column=8).fill = WARN_FILL

# ── Производи→Артикли ───────────────────────────────────────────────────────────────────────────────────────
rows = []
for r in sorted(RECIPES, key=lambda r: (not r["active"], -(r["units_2026"] or 0), r["product_name"])):
    rows.append([r["product_name"], r["kind"], "да" if r["active"] else "не", r["units_2026"], r["units_since_2209"],
                 recipe_text(r["lines"]) or "— нема предлог —", CONF_MK.get(r["confidence"]),
                 "ДА" if r["confidence"] == "high" else "", r["source"], "; ".join(r["evidence"]), r["note"], r["product_id"]])
for e in EXEMPT:
    rows.append([e["product_name"], None, None, None, None, "не е стока", "висока", "ДА", "exempt", "", e["reason"], e["product_id"]])
headers = ["Производ (CRM)", "Вид", "Активен", "Парчиња 2026", "Парчиња од 22.09", "Рецепт: артикли по 1 производ", "Доверба",
           "Автоматски одобри", "Извор", "Докази", "Белешка", "product_id"]
ws, start = sheet(wb, "Производи→Артикли", headers, rows, [40, 9, 8, 11, 11, 60, 10, 10, 22, 60, 40, 38],
                  {4: UNITS, 5: UNITS}, freeze_cols=1, wrap_cols=(1, 6, 10, 11),
                  note="Еден ред = еден CRM производ. Рецептот кажува кои Сигма артикли излегуваат од магацин за 1 продаден производ. "
                       "Висока доверба се одобрува автоматски; средна и ниска ги одобрува сопственикот. Подарок се додава само кога името го кажува.")
for i, row in enumerate(rows):
    c = ws.cell(row=start + i, column=7)
    c.fill = GOOD_FILL if row[6] == "висока" else (WARN_FILL if row[6] in ("ниска", "—") else PatternFill())

# ── collabBox шифри без артикл ──────────────────────────────────────────────────────────────────────────────
STATUS_MK = {"not_stock": "не е стока", "bundle_code": "пакет-шифра → компоненти", "bundle_unresolved": "пакет, не е препознаен",
             "alias_to_sigma": "друга Сигма шифра", "local_article": "нов локален артикл", "name_differs": "името се разликува"}
rows = []
for r in CB:
    rows.append([r["code"], r["role"], r.get("name"), r.get("sigma_name"), r.get("sigma_class"), r.get("name_similarity"),
                 r["lines_since_2209"], r["units_since_2209"], r["docs_since_2209"], STATUS_MK.get(r["status"], r["status"]),
                 r.get("proposal"), CONF_MK.get(r.get("confidence"))])
headers = ["collabBox шифра", "Улога", "Име во collabBox", "Сигма: истата шифра", "Сигма класа", "Сличност на имињата",
           "Ставки од 22.09", "Парчиња од 22.09", "Документи од 22.09", "Статус", "Предлог", "Доверба"]
ws, start = sheet(wb, "collabBox шифри без артикл", headers, rows, [12, 9, 44, 34, 12, 10, 10, 10, 10, 24, 60, 10],
                  {7: UNITS, 8: UNITS3, 9: UNITS}, freeze_cols=1, wrap_cols=(3, 4, 11),
                  note="collabBox ставки од 22.09 чија шифра НЕ е Сигма артикл (или името не се совпаѓа). Сите останати шифри се Сигма шифри и се "
                       "одземаат директно. Слободниот текст во белешките не се прикажува.")
for i, row in enumerate(rows):
    if row[9] in ("нов локален артикл", "пакет, не е препознаен", "името се разликува"):
        ws.cell(row=start + i, column=10).fill = WARN_FILL

# ── Web производи ───────────────────────────────────────────────────────────────────────────────────────────
rows = []
for w in WEB:
    rows.append([w["shop_product_id"], w["variant_id"], w["name"], w.get("variant_label"), ", ".join(w.get("skus") or []),
                 ", ".join(w.get("kinds") or []), w["units_since_2209"], w["lines_all"], (w.get("last_at") or "")[:10],
                 recipe_text(w.get("lines")) or "— нема предлог —", CONF_MK.get(w["confidence"]), w.get("source"),
                 ", ".join(w.get("alias_keys") or [])])
headers = ["Shop ID", "Варијанта", "Име во продавницата (најново)", "Вкус / пакување", "SKU", "Вид", "Парчиња од 22.09",
           "Ставки вкупно", "Последна нарачка", "Артикли", "Доверба", "Извор", "Клучеви"]
ws, start = sheet(wb, "Web производи", headers, rows, [9, 9, 48, 18, 28, 10, 11, 10, 12, 50, 10, 22, 30],
                  {7: UNITS, 8: UNITS}, freeze_cols=1, wrap_cols=(3, 10, 13),
                  note="naturatherapy.mk: веб нарачките немаат collabBox документ, па пратката се одзема по овој предлог (производ / варијанта → Сигма "
                       "артикли). Името е најновото што го има продавницата.")
for i, row in enumerate(rows):
    c = ws.cell(row=start + i, column=11)
    c.fill = GOOD_FILL if row[10] == "висока" else (WARN_FILL if row[10] in ("ниска", "—") else PatternFill())

# ── Пакети ──────────────────────────────────────────────────────────────────────────────────────────────────
KFLAG = {"unpack": "раставување", "last_differs_from_mode": "последното се разликува", "no_whole_assembly": "нема цело склопување",
         "repack": "препакување", "relabel": "преетикетирање", "picked_most_frequent": "земено најчестото"}
rows = []
for k in KITS:
    comp = k["picked"] or k["last_components"]
    rows.append([k["kit_code"], k["kit_name"], recipe_text([dict(code=c, qty=q) for c, q in sorted(comp.items())]),
                 k["assemblies"], k["first"], k["last"], k["picked_doc"] or k["last_doc"],
                 ", ".join(KFLAG.get(f, f) for f in k["flags"]), "да" if k["in_kits_json"] else "не"])
headers = ["Сет шифра", "Сет", "Компоненти за 1 сет", "Склопувања (ММ2)", "Прво", "Последно", "Документ", "Ознаки", "Во kits.json"]
ws, start = sheet(wb, "Пакети", headers, rows, [10, 44, 70, 11, 11, 11, 18, 30, 10], {4: UNITS}, freeze_cols=1, wrap_cols=(2, 3, 8),
                  note="Сигма ММ2 во 04: излез = сетот (ProductionFlag 2), влез = компонентите (минус). ММ2 служи и за преетикетирање и "
                       "препакување — тие не се сетови и не одат во kits.json.")

# ── Сигма документи од 22.09 ────────────────────────────────────────────────────────────────────────────────
VERDICT_MK = {"included": "ВЛЕГУВА", "excluded": "исклучен", "drafted": "најавено"}
rows = []
for d in sorted(DOCS, key=lambda d: (d["doc_date"], d["doc_key"])):
    rows.append([d["doc_key"], d["doc_type"], d.get("type_name"), d["doc_date"], "книжен" if d["status"] == "posted" else "нацрт",
                 (d.get("posted_at") or "")[:16].replace("T", " "), d.get("posted_by"),
                 (d.get("created_at_sigma") or "")[:16].replace("T", " "), d.get("created_by"), d.get("last_change_by"),
                 f"{d['from_obj']} {d.get('from_name') or ''}".strip(), f"{d['to_obj']} {d.get('to_name') or ''}".strip(),
                 d.get("client_code"), d.get("client_name"), d["lines"], d["units_all"],
                 d.get("effect_04") or None, d.get("effect_08") or None, d.get("effect_11") or None, d.get("days_late"),
                 VERDICT_MK.get(d["verdict"], d["verdict"]), d["reason"],
                 "" if d.get("header_in_export") else "заглавието го нема во извозот"])
headers = ["Документ", "Тип", "Тип (име)", "Датум", "Статус", "Книжен на", "Книжел", "Креиран на", "Креирал", "Последна измена",
           "Од", "Кон", "Клиент", "Клиент (фирма)", "Ставки", "Количина (сè)", "Артикли ± во 04", "Артикли ± во 08",
           "Артикли ± во 11", "Книжен по (денови)", "Би влегол?", "Причина / вид", "Забелешка"]
ws, start = sheet(wb, "Сигма документи од 22.09", headers, rows,
                  [20, 7, 22, 11, 9, 16, 10, 16, 10, 10, 30, 30, 9, 28, 8, 12, 11, 11, 11, 10, 11, 40, 22],
                  {15: UNITS, 16: UNITS3, 17: UNITS3, 18: UNITS3, 19: UNITS3, 20: UNITS}, freeze_cols=1, wrap_cols=(3, 11, 12, 14, 22),
                  note="Секој книжен Сигма документ со датум од 22.09 што ги допира 04 / 08 / 11 (+ нацртите). „Би влегол“ = што би направиле "
                       "правилата на серверот: MEX фактурите (000217), АД Астра (000549) и преносите 04↔08 не влегуваат.")
for i, row in enumerate(rows):
    c = ws.cell(row=start + i, column=21)
    c.fill = GOOD_FILL if row[20] == "ВЛЕГУВА" else (WARN_FILL if row[20] == "исклучен" else PatternFill())

path = os.path.join(OUT, f"Magacin-pregled-{ARGS.date}.xlsx")
wb.save(path)
print(f"→ {path}: " + ", ".join(f"{w.title} ({w.max_row})" for w in wb.worksheets))

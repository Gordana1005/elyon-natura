#!/usr/bin/env python3
"""
Owner review file: "Производи надвор од каталогот" (Macedonian CRM, 2026-09-28).

  python scripts/build-uncatalogued-products-xlsx.py exports/products/<date>-analysis.json

Reads the analysis JSON written by scripts/import-catalogue-products.mjs (dry run) and writes
exports/products/<date>-proizvodi-nadvor-od-katalog.xlsx with four sheets:
  Преглед                    what the columns mean + the counts / values + what needs the owner
  Имиња надвор од каталог    one row per normalised name that had sales and no catalogue product
  Нови производи (предлог)   the catalogue rows the script would create
  Каталог (сега)             public.products today, with possible duplicate rows
Money is денари (the Sales tab's allocation of each sale's value over its lines).
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if len(sys.argv) < 2:
    sys.exit("usage: build-uncatalogued-products-xlsx.py <analysis.json>")
src = sys.argv[1]
data = json.load(open(src, encoding="utf-8"))
stamp = os.path.basename(src)[:10]
out = os.path.join(os.path.dirname(src), f"{stamp}-proizvodi-nadvor-od-katalog.xlsx")
rows, cat, new_products, summary = data["rows"], data["catalogue"], data["new_products"], data["summary"]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
CONF_FILL = {
    "exact": PatternFill("solid", fgColor="C6EFCE"),
    "strong": PatternFill("solid", fgColor="E2EFDA"),
    "weak": PatternFill("solid", fgColor="FFEB9C"),
    "none": PatternFill("solid", fgColor="F8CBAD"),
    "new": PatternFill("solid", fgColor="DDEBF7"),
}
KIND_MK = {
    "product": "product — производ", "gift": "gift — подарок", "loyalty_point": "loyalty_point — поени",
    "delivery": "delivery — достава", "note": "note — белешка", "flyer": "flyer — флаер",
}
CONF_MK = {
    "exact": "exact — исто име / иста шифра",
    "strong": "strong — ист производ, друго пакување/пишување",
    "weak": "weak — несигурно, за проверка",
    "none": "none — нема во каталогот",
}
def fmt_int(ws, col, first=2):
    for r in range(first, ws.max_row + 1):
        ws.cell(r, col).number_format = "#,##0"

def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(1, i)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "B2"

def mk(n):
    return f"{round(n):,}".replace(",", ".")

wb = Workbook()

# ── sheet 2: the names ──────────────────────────────────────────────────────
ws = wb.active
ws.title = "Имиња надвор од каталог"
cols = ["Име (најчесто пишување)", "Сите пишувања", "Каде се продавало", "Парчиња", "Нарачки", "Вредност (ден)",
        "Прва продажба", "Последна продажба", "Примери нарачки", "Вид на ставка (предлог)",
        "Предлог од каталогот", "ID на производот (products.id)", "SKU", "Сигурност", "Причина",
        "Предлог акција", "Во скриптата (apply)", "Пакувања по ставка", "Забелешка", "Клуч (alias_norm)"]
header(ws, cols, [46, 50, 34, 10, 10, 14, 13, 13, 28, 22, 40, 38, 13, 30, 70, 30, 12, 11, 30, 40])
for a in rows:
    t = a.get("target") or {}
    new = bool(t.get("newKey"))
    sug = t.get("productName") or t.get("thirdPartyName") or ""
    if new and sug:
        sug = f"{sug} (НОВ)"
    conf = a["confidence"]
    conf_txt = CONF_MK.get(conf, conf)
    if a["kind"] != "product":
        conf_txt = "exact — правило за име (не е производ)"
    elif new and conf == "none":
        conf_txt = "none — нов производ (препознаен)"
    flags = []
    if a.get("bundle"): flags.append("пакет/комбинација")
    if a.get("service"): flags.append("услуга")
    if a.get("thirdParty"): flags.append("трета страна")
    if a.get("giftOnly"): flags.append("само како бесплатна ставка")
    by = a.get("by") or {}
    where = " · ".join(f"{k} ({v['sales']} нар.)" for k, v in sorted(by.items(), key=lambda kv: -kv[1]["value_mkd"]))
    ws.append([
        a["spelling"], " | ".join(a["spellings"]), where, a["units"], a["sales"], a["value_mkd"],
        a["first"], a["last"], ", ".join(a["refs"]), KIND_MK.get(a["kind"], a["kind"]),
        sug, t.get("productId") or "", t.get("sku") or "", conf_txt, a["reason"],
        a.get("action", ""), "да" if a.get("inApply") else "не",
        a.get("multiplier") if a["kind"] == "product" else "", ", ".join(flags), a["norm"],
    ])
    fill = CONF_FILL["new"] if (new and conf == "none") else CONF_FILL.get(conf)
    if fill is not None and a["kind"] == "product":
        ws.cell(ws.max_row, 14).fill = fill
for c in (4, 5, 6):
    fmt_int(ws, c)
for r in range(2, ws.max_row + 1):
    for c in (2, 15):
        ws.cell(r, c).alignment = Alignment(wrap_text=False)
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

# ── sheet 3: new products ───────────────────────────────────────────────────
wn = wb.create_sheet("Нови производи (предлог)")
header(wn, ["Нов производ (име во каталогот)", "Категорија", "Група", "Имиња што ќе се поврзат", "Нарачки",
            "Вредност (ден)", "Прва продажба", "Последна продажба", "ID (детерминистички)"],
       [48, 30, 12, 70, 10, 14, 13, 13, 38])
by_new = {}
for a in rows:
    t = a.get("target") or {}
    if t.get("newKey") and a.get("inApply"):
        by_new.setdefault(t["productId"], []).append(a)
GROUP_MK = {"natura": "Natura", "elixy": "Elixy / I'AM", "device": "уред", "thirdparty": "трета страна"}
def dmy_key(s):
    return (s[6:10] + s[3:5] + s[0:2]) if s else ""
for p in sorted(new_products, key=lambda p: -sum(x["value_mkd"] for x in by_new.get(p["id"], []))):
    names = by_new.get(p["id"], [])
    firsts = [x["first"] for x in names if x["first"]]
    lasts = [x["last"] for x in names if x["last"]]
    wn.append([p["name"], p["category"], GROUP_MK.get(p["group"], p["group"]),
               " | ".join(x["spelling"] for x in names), sum(x["sales"] for x in names),
               sum(x["value_mkd"] for x in names),
               min(firsts, key=dmy_key) if firsts else "", max(lasts, key=dmy_key) if lasts else "", p["id"]])
fmt_int(wn, 5)
fmt_int(wn, 6)
wn.auto_filter.ref = f"A1:I{wn.max_row}"

# ── sheet 4: the catalogue today ────────────────────────────────────────────
wc = wb.create_sheet("Каталог (сега)")
header(wc, ["Име", "SKU", "Активен", "Категорија", "Цена (ден)", "Набавна цена внесена", "Продадено со овој производ (ден, сите години)",
            "Залиха", "Можен дупликат (провери)", "ID"], [44, 13, 9, 18, 11, 12, 18, 9, 60, 38])
for p in sorted(cat, key=lambda p: -p["sold_mkd"]):
    wc.append([p["name"], p["sku"], "да" if p["active"] else "не", p["category"], p["price_mkd"],
               "да" if p["cost"] else "не", p["sold_mkd"], p["stock"], " | ".join(p["duplicates"]), p["id"]])
for c in (5, 7, 8):
    fmt_int(wc, c)
wc.auto_filter.ref = f"A1:J{wc.max_row}"

# ── sheet 1: overview ───────────────────────────────────────────────────────
wo = wb.create_sheet("Преглед", 0)
wo.column_dimensions["A"].width = 44
wo.column_dimensions["B"].width = 20
wo.column_dimensions["C"].width = 20
wo.column_dimensions["D"].width = 90
t = summary["tally"]
def tl(k):
    x = t.get(k, {"names": 0, "value_mkd": 0})
    return x["names"], x["value_mkd"]
sh = summary["shares"]
pct = lambda x: f"{x:.1f} %".replace(".", ",")
lines = [
    ("Производи надвор од каталогот — имиња со продажби", None, None, None),
    (f"Направено {stamp} од scripts/import-catalogue-products.mjs (dry run, run {data['run_id']}). Ништо не е запишано во базата.", None, None, None),
    ("", None, None, None),
    ("Што е ова", None, None, "Секоја продажба (CRM нарачки, веб-продавница, collabBox телешоп) има ставки. Ставка без product_id се води по своето име, "
     "па истиот производ под 20 различни имиња е 20 реда во Увиди → Продажби. Овде е секое такво име (сите години), со предлог кој производ од каталогот е."),
    ("Вредност", None, None, "Вредноста на продажбата (MEX наплата / цена × 61,5 / износ од продавницата) поделена по ставки — истото како табот Продажби. Во денари."),
    ("", None, None, None),
    ("Резиме", "имиња", "вредност (ден)", "што значи"),
]
explain = [
    ("exact", "Исто име по нормализација/транслитерација, или collabBox шифрата на името = SKU во каталогот. Скриптата ги поврзува."),
    ("strong", "Ист производ; разлика само во пакување (1+1, 2+1, 2x, 250/500 ml), вкус, подарок, код за попуст или опис. Скриптата ги поврзува."),
    ("нов производ", "Препознаен производ што го нема во каталогот — скриптата го креира (неактивен, без цена) и ги поврзува имињата."),
    ("трета страна", "Козметика од друг производител продавана на веб (Aura, Wet n Wild, Оливал …) — секој артикл нов производ (може да се исклучи)."),
    ("weak", "Пакет од повеќе платени производи, променета форма, непознат вкус/големина — ЗА ВАС, скриптата не ги допира."),
    ("none", "Именувани пакети без содржина (Terminator, Hulk …), тест ставки — ЗА ВАС."),
    ("услуга", "Јога / фитнес програми — дали да бидат во каталогот? ЗА ВАС."),
    ("вид:note", "ЗАБЕЛЕШКА / упатства за достава во полето за производ — не се производ."),
    ("вид:loyalty_point", "ПОЕН-… — поени за лојалност, не се производ."),
    ("вид:delivery", "ДОСТАВА — трошок за достава, не е производ."),
    ("вид:gift", "Маталка / чашка-мерач — само подарок."),
    ("вид:flyer", "Флаер — не е производ."),
]
for k, why in explain:
    n, v = tl(k)
    lines.append((k, n, v, why))
lines += [
    ("ВКУПНО", summary["names"], summary["value_mkd"], ""),
    ("", None, None, None),
    ("Колку од вредноста на производите е поврзана со каталогот", "пред", "после", ""),
    ("Сите години", pct(sh["all"]["before_pct"]), pct(sh["all"]["after_pct"]),
     f"неповрзано {mk(sh['all']['before_unresolved_mkd'])} → {mk(sh['all']['after_unresolved_mkd'])} ден од {mk(sh['all']['value_mkd'])} ден"),
    ("Септември 2026", pct(sh["sep2026"]["before_pct"]), pct(sh["sep2026"]["after_pct"]),
     f"неповрзано {mk(sh['sep2026']['before_unresolved_mkd'])} → {mk(sh['sep2026']['after_unresolved_mkd'])} ден од {mk(sh['sep2026']['value_mkd'])} ден"),
    ("", None, None, None),
    ("Скриптата (по ваше одобрување)", None, None,
     f"{summary['alias_rows']} алијаси (име → производ / вид) и {summary['new_products']} нови производи "
     f"({summary['new_products_third_party']} од трета страна). Новите производи: неактивни (не се гледаат во формата за нарачка на агентите), "
     "цена 0 и набавна цена 0 — ги внесувате вие; залиха 0."),
    ("", None, None, None),
    ("Колони (лист „Имиња надвор од каталог“)", None, None, None),
    ("Име (најчесто пишување)", None, None, "Најчестото пишување на името во продажбите."),
    ("Сите пишувања", None, None, "Сите пишувања што се сведуваат на исто име (големи/мали букви, празни места)."),
    ("Каде се продавало", None, None, "CRM по извор (AlterCPA / ElyonCRM …), веб-продавница, collabBox телешоп — со број на нарачки."),
    ("Парчиња / Нарачки / Вредност", None, None, "Збир на количини на ставките, број на продажби, вредност во денари."),
    ("Прва / Последна продажба", None, None, "Ден на продажба (dd.mm.yyyy, по Скопје)."),
    ("Примери нарачки", None, None, "До 3 најнови броеви на нарачки (CRM број или број на веб-нарачка)."),
    ("Вид на ставка (предлог)", None, None, "product / gift / loyalty_point / delivery / note / flyer."),
    ("Предлог од каталогот / ID / SKU", None, None, "Постоечки производ (или НОВ) на кој би се поврзало името."),
    ("Сигурност", None, None, "exact · strong · weak · none — види горе."),
    ("Причина", None, None, "Зошто е тој предлог."),
    ("Во скриптата (apply)", None, None, "да = скриптата ќе го поврзе; не = чека ваша одлука."),
    ("Пакувања по ставка", None, None, "Проценка колку пакувања носи една ставка (2+1 → 3). Во извештаите една ставка сè уште се брои по нејзината количина."),
    ("", None, None, None),
    ("Лист „Каталог (сега)“", None, None, "Сите производи денес; „Можен дупликат“ = друг ред што изгледа како ист производ (на пр. Snail Complex и СНАИЛ КОМПЛЕКС cps 30) — за спојување по ваша одлука."),
]
for i, (a, b, c, d) in enumerate(lines, 1):
    wo.cell(i, 1, a)
    if b is not None:
        wo.cell(i, 2, b)
    if c is not None:
        wo.cell(i, 3, c)
    if d is not None:
        wo.cell(i, 4, d)
    wo.cell(i, 4).alignment = Alignment(wrap_text=True, vertical="top")
    wo.cell(i, 1).alignment = Alignment(vertical="top")
    if isinstance(c, (int, float)):
        wo.cell(i, 3).number_format = "#,##0"
wo.cell(1, 1).font = Font(bold=True, size=14)
for i, (a, b, c, d) in enumerate(lines, 1):
    if a in ("Резиме", "Колку од вредноста на производите е поврзана со каталогот", "Колони (лист „Имиња надвор од каталог“)", "ВКУПНО", "Скриптата (по ваше одобрување)"):
        for col in range(1, 5):
            wo.cell(i, col).font = Font(bold=True)

wb.save(out)
print(out)
